import streamlit as st
import pandas as pd
from datetime import datetime, time, timedelta
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage, ImageOps
import io
import os
import json
import base64
import smtplib
from email.message import EmailMessage
import streamlit.components.v1 as components

st.set_page_config(page_title="Control Operacional - Nave Altonorte", layout="wide")

st.title("🏭 Sistema de Control Operacional - Nave Altonorte")
st.markdown("Plataforma web oficial para turnos 4x4.")

# --- ARCHIVO DE RESPALDO EN SERVIDOR ---
RESPALDO_JSON = "respaldo_turno_actual.json"

def guardar_respaldo_disco():
    """Guarda una copia exacta de los datos, supervisores y fotos (en Base64) en disco."""
    try:
        ciclos_serializables = []
        for c in st.session_state.get("ciclos_registrados", []):
            c_copy = c.copy()
            if c_copy.get("Foto_Entrada"):
                c_copy["Foto_Entrada"] = base64.b64encode(c_copy["Foto_Entrada"]).decode('utf-8')
            else:
                c_copy["Foto_Entrada"] = None

            if c_copy.get("Foto_Salida"):
                c_copy["Foto_Salida"] = base64.b64encode(c_copy["Foto_Salida"]).decode('utf-8')
            else:
                c_copy["Foto_Salida"] = None
            
            for campo_t in ["T_Mazamorra", "T_Bloqueo", "T_Ing_Pic", "T_Ret_Pic", "T_Ing_Carg", "T_Ret_Carg", "T_Desbloqueo"]:
                if isinstance(c_copy.get(campo_t), time):
                    c_copy[campo_t] = c_copy[campo_t].strftime("%H:%M:%S")
            ciclos_serializables.append(c_copy)

        datos = {
            "ciclos": ciclos_serializables,
            "termografias": st.session_state.get("termografias_registradas", []),
            "supervisores": {
                "correo_supervisor": st.session_state.get("correo_supervisor", "@glencore.cl"),
                "sup_sop": st.session_state.get("sup_sop", ""),
                "sup_nave": st.session_state.get("sup_nave", ""),
                "sup_crm": st.session_state.get("sup_crm", ""),
                "op_picoton": st.session_state.get("op_picoton", ""),
                "sup_caemin": st.session_state.get("sup_caemin", ""),
                "op_cargador": st.session_state.get("op_cargador", ""),
                "num_regador": st.session_state.get("num_regador", "M-8358")
            },
            "equipos": {
                "est_carg_627": st.session_state.get("est_carg_627", "Revisado por CAEMIN"),
                "est_carg_637": st.session_state.get("est_carg_637", "Pendiente de revisión CAEMIN"),
                "est_aljibe": st.session_state.get("est_aljibe", "Operativo"),
                "est_pic_855": st.session_state.get("est_pic_855", "Revisado por CAEMIN"),
                "est_pic_854": st.session_state.get("est_pic_854", "Pendiente de revisión CAEMIN"),
                "tiempo_teleop": st.session_state.get("tiempo_teleop", 0)
            }
        }
        with open(RESPALDO_JSON, "w", encoding="utf-8") as f:
            json.dump(datos, f)
    except Exception as e:
        print(f"Error al guardar respaldo: {e}")

def cargar_respaldo_disco():
    """Recupera automáticamente los datos desde el disco."""
    if os.path.exists(RESPALDO_JSON):
        try:
            with open(RESPALDO_JSON, "r", encoding="utf-8") as f:
                datos = json.load(f)
                
                ciclos_recuperados = []
                for c in datos.get("ciclos", []):
                    if c.get("Foto_Entrada"):
                        try:
                            c["Foto_Entrada"] = base64.b64decode(c["Foto_Entrada"])
                        except Exception:
                            c["Foto_Entrada"] = None
                    else:
                        c["Foto_Entrada"] = None

                    if c.get("Foto_Salida"):
                        try:
                            c["Foto_Salida"] = base64.b64decode(c["Foto_Salida"])
                        except Exception:
                            c["Foto_Salida"] = None
                    else:
                        c["Foto_Salida"] = None
                    
                    for campo_t in ["T_Mazamorra", "T_Bloqueo", "T_Ing_Pic", "T_Ret_Pic", "T_Ing_Carg", "T_Ret_Carg", "T_Desbloqueo"]:
                        if isinstance(c.get(campo_t), str):
                            c[campo_t] = datetime.strptime(c[campo_t], "%H:%M:%S").time()
                    ciclos_recuperados.append(c)
                    
                st.session_state.ciclos_registrados = ciclos_recuperados
                st.session_state.termografias_registradas = datos.get("termografias", [])
                
                sup_data = datos.get("supervisores", {})
                for k, v in sup_data.items():
                    st.session_state[k] = v
        except Exception as e:
            print(f"Error al cargar respaldo: {e}")

# --- INICIALIZACIÓN DE ESTADO ---
if "ciclos_registrados" not in st.session_state:
    st.session_state.ciclos_registrados = []
if "termografias_registradas" not in st.session_state:
    st.session_state.termografias_registradas = []

if "inicializado" not in st.session_state:
    cargar_respaldo_disco()
    st.session_state.inicializado = True

def sincronizar_localstorage():
    if os.path.exists(RESPALDO_JSON):
        try:
            with open(RESPALDO_JSON, "r", encoding="utf-8") as f:
                contenido_json = f.read()
            js_code = f"""
            <script>
                try {{
                    const dataToSave = {contenido_json};
                    localStorage.setItem("altonorte_respaldo_offline", JSON.stringify(dataToSave));
                }} catch (e) {{
                    console.error("Error guardando en localStorage:", e);
                }}
            </script>
            """
            components.html(js_code, height=0, width=0)
        except Exception:
            pass

sincronizar_localstorage()

def enviar_correo_gmail(destinatario, archivo_bytes, nombre_archivo):
    try:
        remitente = st.secrets.get("GMAIL_USER", "carlos.flores.b@gmail.com")
        password = st.secrets.get("GMAIL_APP_PASSWORD", "eyai squb rdyp unlz")
        
        smtp_server = "smtp.gmail.com"
        smtp_port = 465

        msg = EmailMessage()
        msg['Subject'] = f"📊 Reporte Oficial Control Operacional - {nombre_archivo}"
        msg['From'] = remitente
        msg['To'] = destinatario
        msg.set_content(
            f"Estimado Supervisor (Glencore / Altonorte),\n\n"
            f"Adjunto encontrará la planilla oficial de control operacional de la Nave Altonorte "
            f"correspondiente al turno registrado, incluyendo termografías, ciclos y evidencias fotográficas.\n\n"
            f"Este correo ha sido generado y enviado automáticamente desde el sistema en terreno.\n\n"
            f"Atentamente,\nPlataforma de Control Operacional Altonorte."
        )

        msg.add_attachment(
            archivo_bytes.getvalue(),
            maintype='application',
            subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename=nombre_archivo
        )

        with smtplib.SMTP_SSL(smtp_server, smtp_port) as server:
            server.login(remitente, password)
            server.send_message(msg)
                
        return True, "¡Reporte Excel enviado exitosamente!"
    except Exception as e:
        return False, f"Error al enviar correo: {e}"

def optimizar_imagen(img_bytes, max_width=800, quality=75):
    try:
        img = PILImage.open(io.BytesIO(img_bytes))
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        if img.width > max_width:
            ratio = max_width / float(img.width)
            new_height = int(float(img.height) * ratio)
            img = img.resize((max_width, new_height), PILImage.Resampling.LANCZOS)
        img_con_marco = ImageOps.expand(img, border=6, fill='black')
        buf = io.BytesIO()
        img_con_marco.save(buf, format="JPEG", quality=quality)
        buf.seek(0)
        return buf
    except Exception:
        return io.BytesIO(img_bytes)

def aplicar_borde_cajon_fusionado(ws, start_row, start_col, end_row, end_col, texto, font_titulo):
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
    cell_top_left = ws.cell(row=start_row, column=start_col, value=texto)
    cell_top_left.font = font_titulo
    cell_top_left.alignment = Alignment(horizontal="left", vertical="center")
    medium_side = Side(style='medium', color='000000')
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            top = medium_side if r == start_row else None
            bottom = medium_side if r == end_row else None
            left = medium_side if c == start_col else None
            right = medium_side if c == end_col else None
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)

# --- PANEL SUPERIOR ---
with st.expander("📋 1. Identificación del Turno, Fecha y Supervisores (Toca aquí para desplegar)", expanded=True):
    col_f1, col_f2 = st.columns(2)
    with col_f1:
        fecha_input = st.date_input("Fecha del Turno (Inicio)", datetime.today(), on_change=guardar_respaldo_disco)
    with col_f2:
        tipo_turno = st.selectbox("Tipo de Turno", ["Turno Día (TA)", "Turno Noche (TB)"], on_change=guardar_respaldo_disco)
    
    st.markdown("---")
    st.subheader("Supervisores y Operadores Responsables (Turno 4x4)")
    correo_supervisor = st.text_input("📧 Correo Destinatario (Glencore)", value=st.session_state.get("correo_supervisor", "@glencore.cl"), key="correo_supervisor", on_change=guardar_respaldo_disco)
    
    col_s1, col_s2, col_s3 = st.columns(3)
    with col_s1:
        sup_sop = st.text_input("Supervisor SOP", value=st.session_state.get("sup_sop", ""), key="sup_sop", on_change=guardar_respaldo_disco)
        sup_nave = st.text_input("Supervisor Nave", value=st.session_state.get("sup_nave", ""), key="sup_nave", on_change=guardar_respaldo_disco)
    with col_s2:
        sup_crm = st.text_input("Supervisor CRM", value=st.session_state.get("sup_crm", ""), key="sup_crm", on_change=guardar_respaldo_disco)
        op_picoton = st.text_input("Operador Picotón", value=st.session_state.get("op_picoton", ""), key="op_picoton", on_change=guardar_respaldo_disco)
    with col_s3:
        sup_caemin = st.text_input("Supervisor Caemin", value=st.session_state.get("sup_caemin", ""), key="sup_caemin", on_change=guardar_respaldo_disco)
        op_cargador = st.text_input("Operador Cargador", value=st.session_state.get("op_cargador", ""), key="op_cargador", on_change=guardar_respaldo_disco)

st.markdown("---")

if len(st.session_state.ciclos_registrados) > 0 or len(st.session_state.termografias_registradas) > 0:
    st.success(f"🛡️ **Modo Alta Confiabilidad Activo:** `{len(st.session_state.ciclos_registrados)}` ciclos y `{len(st.session_state.termografias_registradas)}` termografías respaldadas.")

tab1, tab2, tab3, tab4 = st.tabs([
    "1. Checklist & Equipos", 
    "2. Termografía (°C)", 
    "3. Ciclos CPS & Fotos", 
    "4. Consolidado y Envío correo"
])

with tab1:
    st.subheader("Control Previo de Inicio de Actividades y Equipos")
    col_e1, col_e2 = st.columns(2)
    with col_e1:
        est_carg_627 = st.selectbox("Cargador M-627", ["Revisado por CAEMIN", "Pendiente de revisión CAEMIN", "Fuera de servicio"], key="est_carg_627", on_change=guardar_respaldo_disco)
        est_carg_637 = st.selectbox("Cargador M-637", ["Pendiente de revisión CAEMIN", "Revisado por CAEMIN", "Fuera de servicio"], key="est_carg_637", on_change=guardar_respaldo_disco)
        num_regador = st.text_input("🚰 Número de Regador / Aljibe", value=st.session_state.get("num_regador", "M-8358"), key="num_regador", on_change=guardar_respaldo_disco)
        est_aljibe = st.selectbox(f"Camión Aljibe ({num_regador})", ["Operativo", "No disponible", "Fuera de servicio"], key="est_aljibe", on_change=guardar_respaldo_disco)
    with col_e2:
        est_pic_855 = st.selectbox("Picotón M-855", ["Revisado por CAEMIN", "Pendiente de revisión CAEMIN", "Fuera de servicio"], key="est_pic_855", on_change=guardar_respaldo_disco)
        est_pic_854 = st.selectbox("Picotón M-854", ["Pendiente de revisión CAEMIN", "Revisado por CAEMIN", "Fuera de servicio"], key="est_pic_854", on_change=guardar_respaldo_disco)
    
    tiempo_teleop = st.number_input("⏱️ Tiempo en teleoperación (minutos perdidos por señal)", min_value=0, value=st.session_state.get("tiempo_teleop", 0), key="tiempo_teleop", on_change=guardar_respaldo_disco)

with tab2:
    st.subheader("🌡️ Registro de Termografía (°C) por Ingreso")
    col_t_eq, col_t_cps = st.columns(2)
    with col_t_eq:
        eq_term = st.selectbox("Equipo", ["Picotón", "Cargador"], key="eq_term")
    with col_t_cps:
        cps_term = st.selectbox("CPS", ["CPS-1", "CPS-2", "CPS-3", "CPS-4"], key="term_cps")
        
    col_val1, col_val2, col_val3 = st.columns(3)
    if eq_term == "Picotón":
        with col_val1: val1 = st.number_input("Flexibles (°C)", value=40.0, key="t_flex")
        with col_val2: val2 = st.number_input("Cuña (°C)", value=160.0, key="t_cuna")
        with col_val3: val3 = st.number_input("Tornamesa (°C)", value=20.0, key="t_torna")
        p_puntos = {"Flexibles": val1, "Cuña": val2, "Tornamesa": val3}
    else:
        with col_val1: val1 = st.number_input("Batea (°C)", value=140.0, key="t_batea")
        with col_val2: val2 = st.number_input("Parte inferior (°C)", value=200.0, key="t_inf")
        with col_val3: val3 = st.number_input("Flexibles (°C)", value=40.0, key="t_flex_carg")
        p_puntos = {"Batea": val1, "Parte Inferior": val2, "Flexibles": val3}
        
    if st.button("➕ Agregar Registro de Termografía"):
        st.session_state.termografias_registradas.append({
            "Equipo": eq_term,
            "CPS": cps_term,
            **p_puntos,
            "Hora": datetime.now().strftime("%H:%M")
        })
        guardar_respaldo_disco()
        st.success("¡Termografía agregada y respaldada con éxito!")
        st.rerun()
        
    if len(st.session_state.termografias_registradas) > 0:
        st.markdown(f"### Historial de Termografías (Total: {len(st.session_state.termografias_registradas)})")
        df_term = pd.DataFrame(st.session_state.termografias_registradas)
        df_term.index = range(1, len(df_term) + 1)
        st.dataframe(df_term, use_container_width=True)

with tab3:
    st.subheader("🔄 Registro de Ciclos CPS, Horarios y Desviaciones")
    col_sel1, col_sel2 = st.columns(2)
    with col_sel1:
        cps_seleccionado = st.selectbox("Seleccione CPS", ["CPS-1", "CPS-2", "CPS-3", "CPS-4"], key="ciclo_cps")
    with col_sel2:
        nombre_ciclo = st.text_input("Identificador de Ciclo (ej. J-117, J-118)", "J-117", key="nombre_ciclo_input")
    
    col_h1, col_h2 = st.columns(2)
    with col_h1:
        fecha_ciclo_inicio = st.date_input("Fecha Inicio Ciclo", fecha_input, key="f_ciclo_inicio")
        t_mazamorra = st.time_input("1. Término retiro mazamorra desde CPS", value=time(20, 25), key="t_maz")
        t_bloqueo = st.time_input("2. Bloqueo de tapas CPS", value=time(20, 31), key="t_bloq")
        t_ing_pic = st.time_input("3. Ingreso Picotón foso", value=time(20, 32), key="t_ipic")
        t_ret_pic = st.time_input("4. Retiro Picotón foso", value=time(20, 40), key="t_rpic")
    with col_h2:
        t_ing_carg = st.time_input("5. Ingreso Cargador Foso", value=time(20, 42), key="t_icarg")
        t_ret_carg = st.time_input("6. Retiro de Cargador en foso", value=time(20, 48), key="t_rcarg")
        t_desbloqueo = st.time_input("7. Desbloqueo de tapas CPS", value=time(20, 53), key="t_desb")

    col_obs1, col_obs2 = st.columns(2)
    with col_obs1:
        obs_mazamorra = st.text_input("Nota Retiro Mazamorra", value="", key="obs_maz")
        obs_bloqueo = st.text_input("Nota Bloqueo de Tapas", value="", key="obs_bloq")
        obs_ing_pic = st.text_input("Nota Ingreso Picotón", value="", key="obs_ipic")
        obs_ret_pic = st.text_input("Nota Retiro Picotón", value="", key="obs_rpic")
    with col_obs2:
        obs_ing_carg = st.text_input("Nota Ingreso Cargador", value="", key="obs_icarg")
        obs_ret_carg = st.text_input("Nota Retiro Cargador", value="", key="obs_rcarg")
        obs_desbloq = st.text_input("Nota Desbloqueo de Tapas", value="", key="obs_desb")

    col_val1, col_val2, col_val3 = st.columns(3)
    with col_val1: num_baldadas = st.number_input("Número de Baldadas", min_value=0, value=2, key="num_bal")
    with col_val2: toneladas_ext = st.number_input("Toneladas estimadas", value=8.0, key="ton_ext")
    with col_val3: temp_mat_ret = st.number_input("Temperatura material retirado (°C)", value=545.7, key="temp_mat")

    col_f1, col_f2 = st.columns(2)
    with col_f1:
        foto_entrada = st.file_uploader(f"Foto ANTES del Ingreso ({cps_seleccionado})", type=["jpg", "jpeg", "png"], key=f"f_ent_{nombre_ciclo}")
    with col_f2:
        foto_salida = st.file_uploader(f"Foto a la SALIDA ({cps_seleccionado})", type=["jpg", "jpeg", "png"], key=f"f_sal_{nombre_ciclo}")

    if st.button("➕ Guardar Ciclo y Sus Fotografías"):
        b_ent = foto_entrada.getvalue() if foto_entrada else None
        b_sal = foto_salida.getvalue() if foto_salida else None

        st.session_state.ciclos_registrados.append({
            "CPS": cps_seleccionado,
            "Ciclo": nombre_ciclo,
            "Fecha_Inicio": fecha_ciclo_inicio,
            "T_Mazamorra": t_mazamorra,
            "T_Bloqueo": t_bloqueo,
            "T_Ing_Pic": t_ing_pic,
            "T_Ret_Pic": t_ret_pic,
            "T_Ing_Carg": t_ing_carg,
            "T_Ret_Carg": t_ret_carg,
            "T_Desbloqueo": t_desbloqueo,
            "Obs_Mazamorra": obs_mazamorra,
            "Obs_Bloqueo": obs_bloqueo,
            "Obs_Ing_Pic": obs_ing_pic,
            "Obs_Ret_Pic": obs_ret_pic,
            "Obs_Ing_Carg": obs_ing_carg,
            "Obs_Ret_Carg": obs_ret_carg,
            "Obs_Desbloqueo": obs_desbloq,
            "Baldadas": num_baldadas,
            "Toneladas": toneladas_ext,
            "Temp_Material": temp_mat_ret,
            "Foto_Entrada": b_ent,
            "Foto_Salida": b_sal
        })
        guardar_respaldo_disco()
        st.success(f"¡Ciclo {nombre_ciclo} guardado con éxito!")
        st.rerun()

with tab4:
    st.subheader("📋 Consolidado y Envío por correo")
    
    # LÓGICA BLINDADA DE FECHAS OFICIALES PARA TURNO NOCHE (TB) vs DÍA (TA)
    if "Noche" in tipo_turno or "TB" in tipo_turno:
        sufijo_turno = "TB"
        # Si es Turno Noche que inicia el 27, el turno oficial de cierre es el 27 (o según el criterio exacto de fecha de reporte)
        # Tu observación anterior: "ayer fue el turno del día 27 por que aparece como 28?"
        # Si inicias el turno noche el 27 a las 20:00, la fecha oficial del reporte de ese turno es el 27 TB.
        fecha_turno_oficial = fecha_input
    else:
        sufijo_turno = "TA"
        fecha_turno_oficial = fecha_input

    str_fecha = fecha_turno_oficial.strftime("%d-%m")
    nombre_nueva_hoja = f"{str_fecha} {sufijo_turno}"
    nombre_archivo_salida = f"Control_Nave_{str_fecha}_{sufijo_turno}.xlsx"

    st.info(f"Correo Destinatario (Glencore): **{correo_supervisor}** | Turno Oficial: **{str_fecha} {sufijo_turno}**")
    
    if st.button("🗑️ Finalizar y Limpiar Turno Actual"):
        st.session_state.ciclos_registrados = []
        st.session_state.termografias_registradas = []
        if os.path.exists(RESPALDO_JSON):
            os.remove(RESPALDO_JSON)
        components.html("<script>localStorage.removeItem('altonorte_respaldo_offline');</script>", height=0, width=0)
        st.success("Turno finalizado y memoria limpiada.")
        st.rerun()

    st.markdown("---")
    if st.button("📧 Generar y Enviar Planilla Excel por correo"):
        base_excel = 'plantilla.xlsx'
        
        try:
            if os.path.exists(base_excel):
                wb = openpyxl.load_workbook(base_excel)
                
                # SELECCIÓN ESTRICTA DE LA PLANILLA Válida (Busca la hoja más reciente formateada, ej. '28-10 TB' o la última hoja activa)
                hoja_nombre_origen = wb.sheetnames[-1] # Toma la última hoja formateada existente como base
                ws_source = wb[hoja_nombre_origen]
                
                if nombre_nueva_hoja in wb.sheetnames:
                    del wb[nombre_nueva_hoja]
                
                ws = wb.copy_worksheet(ws_source)
                ws.title = nombre_nueva_hoja
            else:
                st.error("No se encontró el archivo 'plantilla.xlsx' en el repositorio.")
                st.stop()

            # Inyección de cabeceras y fechas correctas
            ws['B3'] = f"Fecha: {str_fecha} {sufijo_turno}"
            ws['C4'] = sup_sop
            ws['E4'] = sup_crm
            ws['F4'] = "Operador Picotón: "
            ws['G4'] = op_picoton
            
            ws['C5'] = sup_caemin
            ws['E5'] = sup_nave
            ws['F5'] = "Operador Cargador: "
            ws['G5'] = op_cargador

            tick_symbol = "✓"
            for r_chk in range(8, 15):
                cell_chk = ws.cell(row=r_chk, column=5, value=tick_symbol)
                cell_chk.alignment = Alignment(horizontal="center", vertical="center")

            comentarios_generales = (
                f"Operativos\n"
                f"Cargador M-627 {est_carg_627}. \n"
                f"Picoton M-855 {est_pic_855}. \n"
                f"Camión Aljibe ({num_regador}) {est_aljibe}.\n\n"
                f"Stand by:\n"
                f"Picoton 854 {est_pic_854}.\n"
                f"Cargador 637 {est_carg_637}.\n\n"
                f"Tiempo en teleoperación: {tiempo_teleop} min por perdida de señal."
            )
            ws.cell(row=8, column=6, value=comentarios_generales)
            ws['F14'] = f"Camión Aljibe {num_regador} {est_aljibe.lower()}"

            p_pic_rows = [18, 19, 20]
            c_car_rows = [23, 24, 25]
            
            idx_p = 0
            idx_c = 0
            for term in st.session_state.termografias_registradas:
                if term["Equipo"] == "Picotón" and idx_p < len(p_pic_rows):
                    r = p_pic_rows[idx_p]
                    ws.cell(row=r, column=2, value=term["CPS"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=3, value=term["Flexibles"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=4, value=term["Cuña"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=5, value=term["Tornamesa"]).alignment = Alignment(horizontal="center", vertical="center")
                    idx_p += 1
                elif term["Equipo"] == "Cargador" and idx_c < len(c_car_rows):
                    r = c_car_rows[idx_c]
                    ws.cell(row=r, column=2, value=term["CPS"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=3, value=term["Batea"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=4, value=term["Parte Inferior"]).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=r, column=5, value=term["Flexibles"]).alignment = Alignment(horizontal="center", vertical="center")
                    idx_c += 1

            for col_l in ['B', 'C', 'D', 'E', 'F', 'G']:
                for row_l in range(63, 105):
                    cell_target = ws[f'{col_l}{row_l}']
                    if cell_target.value and ("CPS" in str(cell_target.value) or "Antes" in str(cell_target.value) or "Despues" in str(cell_target.value)):
                        cell_target.value = None

            ciclo_start_rows = [28, 40, 52]
            image_target_rows = [67, 81, 98]

            font_titulo = Font(name="Aptos Narrow", size=16, bold=True, color="000000")

            for idx, ciclo in enumerate(st.session_state.ciclos_registrados):
                if idx >= len(ciclo_start_rows):
                    break
                
                r = ciclo_start_rows[idx]
                img_row = image_target_rows[idx]
                
                ws.cell(row=r, column=3, value=ciclo["CPS"])
                ws.cell(row=r+1, column=3, value=ciclo["Ciclo"])
                
                def parse_time(t_val):
                    if isinstance(t_val, str):
                        return datetime.strptime(t_val, "%H:%M:%S").time()
                    return t_val

                t_maz = parse_time(ciclo["T_Mazamorra"])
                t_blo = parse_time(ciclo["T_Bloqueo"])
                t_i_pic = parse_time(ciclo["T_Ing_Pic"])
                t_r_pic = parse_time(ciclo["T_Ret_Pic"])
                t_i_car = parse_time(ciclo["T_Ing_Carg"])
                t_r_car = parse_time(ciclo["T_Ret_Carg"])
                t_des = parse_time(ciclo["T_Desbloqueo"])

                f_base = ciclo.get("Fecha_Inicio", fecha_input)
                if isinstance(f_base, str):
                    f_base = datetime.strptime(f_base, "%Y-%m-%d").date()

                # Lógica estricta de incremento de día solo si pasa la medianoche real
                d_maz = f_base
                d_blo = d_maz if t_blo >= t_maz else d_maz + timedelta(days=1)
                d_i_pic = d_blo if t_i_pic >= t_blo else d_blo + timedelta(days=1)
                d_r_pic = d_i_pic if t_r_pic >= t_i_pic else d_i_pic + timedelta(days=1)
                d_i_car = d_r_pic if t_i_car >= t_r_pic else d_r_pic + timedelta(days=1)
                d_r_car = d_i_car if t_r_car >= t_i_car else d_i_car + timedelta(days=1)
                d_des = d_r_car if t_des >= t_r_car else d_r_car + timedelta(days=1)

                hitos = [
                    (datetime.combine(d_maz, t_maz), r+3),
                    (datetime.combine(d_blo, t_blo), r+4),
                    (datetime.combine(d_i_pic, t_i_pic), r+5),
                    (datetime.combine(d_r_pic, t_r_pic), r+6),
                    (datetime.combine(d_i_car, t_i_car), r+7),
                    (datetime.combine(d_r_car, t_r_car), r+8),
                    (datetime.combine(d_des, t_des), r+9)
                ]
                
                for h_dt, h_row in hitos:
                    cell_h = ws.cell(row=h_row, column=3, value=h_dt)
                    cell_h.number_format = 'DD-MM-YYYY HH:MM:SS'
                
                observaciones_hitos = [
                    (ciclo.get("Obs_Mazamorra"), r+3),
                    (ciclo.get("Obs_Bloqueo"), r+4),
                    (ciclo.get("Obs_Ing_Pic"), r+5),
                    (ciclo.get("Obs_Ret_Pic"), r+6),
                    (ciclo.get("Obs_Ing_Carg"), r+7),
                    (ciclo.get("Obs_Ret_Carg"), r+8),
                    (ciclo.get("Obs_Desbloqueo"), r+9)
                ]
                for obs_text, obs_row in observaciones_hitos:
                    if obs_text:
                        ws.cell(row=obs_row, column=6, value=obs_text)

                row_ing_carg = r + 7
                if ciclo.get("Temp_Material") is not None:
                    ws.cell(row=row_ing_carg, column=6, value="Temperatura material retirado (°C)")
                    ws.cell(row=row_ing_carg, column=7, value=str(ciclo["Temp_Material"]))

                ret_carg_row = r + 8
                ws.cell(row=ret_carg_row, column=5, value=ciclo["Baldadas"])
                ws.cell(row=ret_carg_row, column=6, value=ciclo["Toneladas"])
                ws.cell(row=ret_carg_row, column=7, value="Ton. ")

                titulo_row = img_row - 1

                if ciclo.get("Foto_Entrada") is not None:
                    try:
                        buf_ent = optimizar_imagen(ciclo["Foto_Entrada"])
                        txt_ent = f"📸 FOTO ANTES DE LIMPIEZA - {ciclo['CPS']} ({ciclo['Ciclo']})"
                        aplicar_borde_cajon_fusionado(ws, start_row=titulo_row, start_col=2, end_row=titulo_row, end_col=3, texto=txt_ent, font_titulo=font_titulo)
                        img_ent = XLImage(buf_ent)
                        img_ent.width = 180
                        img_ent.height = 140
                        ws.add_image(img_ent, f"B{img_row}")
                    except Exception as img_err:
                        st.warning(f"Error foto entrada {ciclo['CPS']}: {img_err}")

                if ciclo.get("Foto_Salida") is not None:
                    try:
                        buf_sal = optimizar_imagen(ciclo["Foto_Salida"])
                        txt_sal = f"📸 FOTO DESPUES DE LIMPIEZA - {ciclo['CPS']} ({ciclo['Ciclo']})"
                        aplicar_borde_cajon_fusionado(ws, start_row=titulo_row, start_col=5, end_row=titulo_row, end_col=7, texto=txt_sal, font_titulo=font_titulo)
                        img_sal = XLImage(buf_sal)
                        img_sal.width = 180
                        img_sal.height = 140
                        ws.add_image(img_sal, f"E{img_row}")
                    except Exception as img_err:
                        st.warning(f"Error foto salida {ciclo['CPS']}: {img_err}")

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            exito, mensaje = enviar_correo_gmail(correo_supervisor, output, nombre_archivo_salida)
            if exito:
                st.success(mensaje)
            else:
                st.error(mensaje)
                st.download_button(
                    label="📥 Descarga Alternativa de Planilla Excel",
                    data=output,
                    file_name=nombre_archivo_salida,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
        except Exception as e:
            st.error(f"Error al generar o enviar el archivo: {e}")
